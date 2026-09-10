"""judge 标注材料提取（纯函数）：从 judge generation 的输入 prompt 切出人审材料。

judge 调用经 gateway 观测，观测 input 落库为
{"messages": [{"role": "user", "content": <渲染后 rubric>}]}；渲染后 rubric 以
【小节】内嵌该维度 judge 实际看到的关键变量（【用户查询】/【分析报告】/
【多空辩论记录】/【交易决策】等，见 evals/judges.py RUBRICS）。

本模块按该维度 rubric 的**已知小节标记**定位切分（rubric 正文嵌套的未知
【】标记如【technical】【bull】保留在材料原文，不当作小节被切断），尾部用
「只输出 JSON / 评估」等 rubric 指令锚点截掉，落进 CSV 标注表——judge 看到
什么，标注人就审什么（同口径）。
"""

from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path
from typing import Any

_SPACE_RE = re.compile(r"\s+")

# 切分后做尾部锚点：rubric 指令/评分标准从这些短语开始，截掉避免混进材料
_TAIL_ANCHORS = ("只输出 JSON", "\n评估", "\n若交易决策含", "\n无 evidence_refs")

# 各维度人审需要的小节（与 judges.RUBRICS 的 {{var}} 对应【小节名】对齐；
# 新增/改名小节需同步此处并跑测试）
DIMENSION_SECTIONS: dict[str, tuple[str, ...]] = {
    "report_relevance": ("【用户查询】", "【分析报告】"),
    "debate_quality": ("【多空辩论记录】",),
    "decision_grounding": (
        "【分析师结论】",
        "【多空辩论记录】",
        "【Research Manager 结论】",
        "【交易决策】",
    ),
    "consistency": (
        "【分析师章节结论】",
        "【Research Manager 结论】",
        "【Risk Judge 裁决】",
        "【Fund Manager 最终决策】",
        "【最终报告结论章节】",
    ),
}

# rubric 首句特征词 → 维度（judge generation 观测名统一为 "judge"，须按此区分）
RUBRIC_INTRO: dict[str, str] = {
    "report_relevance": "投资研究报告评审专家",
    "debate_quality": "投资辩论质量评审专家",
    "decision_grounding": "投资决策依据评审专家",
    "consistency": "投资报告一致性评审专家",
}

CSV_HEADER = [
    "trace_id",
    "dimension",
    "judge_score",
    "judge_reason",
    "human_score",
    "confidence",
    "trace_url",
    "material_summary",
]

# 盲标版表头：不含 LLM judge 分/理由（标注人先独立打分，防锚定效应；
# judge 分由 --blind-judge jsonl 独立存储，measure 合并计算一致性）
BLIND_HEADER = [
    "trace_id",
    "dimension",
    "human_score",
    "confidence",
    "trace_url",
    "material_summary",
]

# 标注伴读文件：rubric 速查 + 指标速查（只求识别与数量级直觉，不做计算）
METRIC_CHEATSHEET = """## 指标速查（标注用：只求识别与数量级直觉，不做计算；数字对错由校验器负责）

- ROE（净资产收益率）：股东投入资本一年的回报率；≈10%+ 属不错，>20% 优秀，<5% 偏弱。
- 毛利率：收入扣除直接成本后的留存；>40% 常见于白酒/医药/软件，<10% 属薄利制造。
- 净利率：最终净利占收入比例；通常明显低于毛利率。
- PE（市盈率）：股价 ÷ 每股盈利；<15 偏便宜/周期，>40 高预期/成长。
- PB（市净率）：股价 ÷ 每股净资产；>1 有溢价，<1 破净。
- 经营现金流：生意实际收到的钱；持续为正才是真赚钱。
- 资产负债率：总负债 ÷ 总资产；>70% 杠杆偏高须警惕。
- 回撤/波动率：回撤大 = 曾深度亏损；波动率大 = 价格上蹿下跳。
"""


def sections_for_dimension(dimension: str) -> tuple[str, ...]:
    """该维度人审需要的【小节】名列表（未知维度返回空元组）。"""
    return DIMENSION_SECTIONS.get(dimension, ())


def detect_dimension(rendered: str | None) -> str | None:
    """按 rubric 首句特征识别渲染 prompt 所属维度（无法识别返回 None）。"""
    if not rendered:
        return None
    for dim, intro in RUBRIC_INTRO.items():
        if intro in rendered:
            return dim
    return None


def _clean(text: str) -> str:
    """压缩空白（含换行）为单个空格，摘要紧凑。"""
    return " ".join(_SPACE_RE.split(text or "")).strip()


def extract_sections(rendered: str, dimension: str) -> dict[str, str]:
    """按该维度 rubric 的已知小节标记定位切分 → {小节名: 文本}。

    只以 DIMENSION_SECTIONS 中的标记为边界（内容里的嵌套未知【】标记不切分）；
    每个小节的文本取「本标记之后 → 下一已知标记 / 尾部锚点之前」。rubric 模板
    标记始终出现（{{var}} 为空时该节文本为空串，build_summary 跳过）。
    """
    markers = sections_for_dimension(dimension)
    if not rendered or not markers:
        return {}
    positions: list[tuple[int, str]] = []
    for m in markers:
        idx = rendered.find(m)
        if idx >= 0:
            positions.append((idx, m))
    positions.sort()
    out: dict[str, str] = {}
    for k, (idx, m) in enumerate(positions):
        start = idx + len(m)
        end = positions[k + 1][0] if k + 1 < len(positions) else _tail_start(rendered, start)
        out[m] = _clean(rendered[start:end])
    return out


def _tail_start(rendered: str, start: int) -> int:
    """末节文本的尾部锚点：rubric 指令/评分标准起始位置（找不到返回全文末尾）。"""
    hits = [rendered.find(anchor, start) for anchor in _TAIL_ANCHORS]
    hits = [h for h in hits if h >= 0]
    return min(hits) if hits else len(rendered)


def prompt_from_observation_input(raw: Any) -> str | None:
    """观测 input 各形态 → 渲染 prompt 文本；取不到返回 None（不崩）。

    gateway 落库形态：{"messages": [{"role": "user", "content": <prompt>}]}；
    兼容直传 {"content": ...}、JSON 字符串、裸文本。
    """
    if raw is None:
        return None
    if isinstance(raw, str):
        s = raw.strip()
        if not s:
            return None
        try:
            parsed = json.loads(s)
        except json.JSONDecodeError:
            return s  # 裸文本即 prompt
        return prompt_from_observation_input(parsed)
    if isinstance(raw, dict):
        messages = raw.get("messages")
        if messages is not None:
            return prompt_from_observation_input(messages)
        content = raw.get("content")
        return content if isinstance(content, str) and content else None
    if isinstance(raw, list):
        for item in raw:
            got = prompt_from_observation_input(item)
            if got:
                return got
    return None


# 各维度人审摘要的每节截断上限：judge 变量最小预算 4096 字节（≈1300 中文字），
# 展示层 5000 字符可完整覆盖任一变量——**统一 5000，不留维度分层**。历史教训：
# 分维度默认 350 时 report_relevance 的【分析报告】被腰斩（实测 acb17607：材料在
# 表格中间截断「机器人叙…」，judge 看到的完整报告含「一句话总结」结论，标注人
# 只见到一半——1 vs 5 的假分歧）；consistency 的 RM 结论同样被砍（81a133c2）。
# 材料必须完整呈现 judge 所见，截断无正当场景。
_DEFAULT_SUMMARY_LIMIT = 5000


# 分析师类小节：其文本是各 agent 自带 summary 的拼贴（extract._summarize_analyst_reports
# 以 【technical】/【macro】/【fundamental】/【sentiment】 标记每份），按 agent 分行完整
# 展示——agent 的总结性摘要在数据里就有（report.summary），不应退化为截取式预览。
_AGENT_SECTIONS = ("【分析师章节结论】", "【分析师结论】")
_AGENT_MARKER_RE = re.compile(r"【([^】]+)】")


def _linebreak_agent_sections(text: str) -> str:
    """把分析师节原文按【agent名】标记插入换行（内容不变，只加分隔）。

    原文是 extract 拼贴的密排段（4 个 agent 连成一坨无法分辨边界）；「▼ 各
    agent 摘要」预览与其完全重复、无信息量（2026-09-09 实测），已移除——agent
    节只保留原文，按 agent 分行展示。
    """
    if "【" not in text:
        return text.strip()
    return _AGENT_MARKER_RE.sub(r"\n【\1】", text).strip()


def build_summary(rendered: str | None, dimension: str, *, limit: int | None = None) -> str:
    """按维度拼人审摘要：小节标题 + 完整原文。

    分析师类小节按【agent】标记分行展示原文（与 judge 所见同内容，仅加换行
    分隔）；「▼ 各 agent 摘要」预览与原文完全重复，已移除（2026-09-09）。
    默认 limit 为 _DEFAULT_SUMMARY_LIMIT(5000)——judge 变量最小预算 4096 字节，
    展示层完整呈现、不分维度分层（历史教训见常量注释）。显式传 limit 时以显式值为准。
    agent 节是各 agent 自带总结性摘要（数据里就有），
    不应用 limit 二次截短——350 只够排第一的 technical 开头，其余 agent 结论会被切掉
    （2026-09-09 回归：consistency 维度标注人看不到基本面/情绪面结论，无法核对一致性）。
    无渲染 prompt / 小节缺失 → 返回「材料缺失」提示（标注人点 trace_url 查看）。
    """
    if not rendered:
        return "材料缺失：请点 trace_url 在 Langfuse 中查看该维度对应内容"
    limit = limit if limit is not None else _DEFAULT_SUMMARY_LIMIT
    sections = extract_sections(rendered, dimension)
    parts: list[str] = []
    for label in sections_for_dimension(dimension):
        text = sections.get(label, "")
        if not text:
            continue
        # 人读渲染：决策类 JSON（交易决策/风控裁决，含 "" 双重转义）→ 人读格式；
        # 信息等价仅格式变化，解析失败保持原文。须在 limit 截断前执行（截断会
        # 切断 JSON 块导致解析失败）。
        text = humanize_json_blocks(text)
        if label not in _AGENT_SECTIONS and len(text) > limit:
            text = text[:limit] + "…"
        if label in _AGENT_SECTIONS and "【" in text:
            text = _linebreak_agent_sections(text)
        parts.append(f"{label}\n{text}")
    if not parts:
        return "材料缺失：渲染 prompt 中未找到该维度小节（请点 trace_url 查看）"
    return "\n".join(parts)


_JSON_BLOCK_RE = re.compile(r"\{.*\}", re.DOTALL)

# 决策类 JSON 的人读字段顺序（出现的渲染，未出现的跳过）
_DECISION_FIELDS: tuple[tuple[str, str], ...] = (
    ("action", "action"),
    ("decision", "decision"),
    ("confidence", "置信度"),
    ("position_size", "仓位"),
    ("entry_price", "入场价"),
    ("stop_loss", "止损"),
    ("target_price", "目标价"),
)


def _humanize_decision_obj(obj: dict) -> str:
    """决策 dict → 人读多行文本（信息等价，仅格式变化）。"""
    lines: list[str] = []
    for key, label in _DECISION_FIELDS:
        val = obj.get(key)
        if val in (None, ""):
            continue
        if key == "confidence":
            lines.append(f"{label}: {float(val):.2f}")
        else:
            lines.append(f"{label}: {val}")
    refs = obj.get("evidence_refs") or []
    if isinstance(refs, list) and refs:
        from collections import Counter

        sources = Counter(str((r or {}).get("source", "?")) for r in refs if isinstance(r, dict))
        dist = ", ".join(f"{s}×{n}" for s, n in sources.most_common())
        lines.append(f"论据引用 {len(refs)} 条（{dist}）")
    reasoning = str(obj.get("reasoning") or "").strip()
    if reasoning:
        lines.append(f"理由: {reasoning}")
    return "\n".join(lines) if lines else None


def humanize_json_blocks(text: str) -> str:
    """把材料文本中的决策类 JSON 块（含 "" 双重转义形态）渲染为人读格式。

    信息内容等价（action/置信度/仓位/论据来源分布/理由分行），仅格式变化——
    标注人读的是 judge 所见的同一信息，不再啃转义 JSON（delta 3.7）。
    解析失败保持原文，绝不丢内容。
    """
    if not text or "{" not in text:
        return text
    m = _JSON_BLOCK_RE.search(text)
    if not m:
        return text
    raw = m.group(0)
    try:
        obj = json.loads(raw.replace('""', '"'))
    except (json.JSONDecodeError, ValueError):
        return text
    if not isinstance(obj, dict):
        return text
    rendered = _humanize_decision_obj(obj)
    if not rendered:
        return text
    return text.replace(raw, "\n" + rendered + "\n")


def to_csv(rows: list[dict[str, Any]]) -> str:
    """标注表 → CSV 文本（含表头；缺列补空串，human_score/confidence 待人工回填）。"""
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(CSV_HEADER)
    for row in rows:
        writer.writerow([str(row.get(k) or "") for k in CSV_HEADER])
    return buf.getvalue()


# ── xlsx 标注表（Excel 直开免拉格子：冻结表头 / 列宽 / 自动换行 / 行高估算）──

_XLSX_COL_WIDTHS = {
    "trace_id": 36,
    "dimension": 17,
    "judge_score": 10,
    "judge_reason": 46,
    "human_score": 11,
    "confidence": 10,
    "trace_url": 44,
    "material_summary": 96,
}
_HEADER_FILL = "D9E1F2"
_FILL_COLUMNS = ("human_score", "confidence")  # 待人工回填列高亮


def _estimate_row_height(text: str, width_chars: int) -> int:
    """按中文字符约占 2 个宽度单位估算换行行数 → 行高（pt），上限防撑爆。"""
    if not text:
        return 22
    eff = max(4, width_chars // 2)
    lines = sum(1 + len(seg) // eff for seg in text.split("\n"))
    return min(max(lines * 13, 22), 320)


def to_xlsx(
    rows: list[dict[str, Any]], path: str | Path, *, columns: list[str] | None = None
) -> None:
    """标注表 → .xlsx（Excel 直开的成品：冻结表头、每列宽、长文本自动换行、
    行高按内容估算、trace_url 可点击、human_score/confidence 列高亮提醒回填）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cols = columns or CSV_HEADER
    wb = Workbook()
    ws = wb.active
    ws.title = "标注表"
    for col, name in enumerate(cols, 1):
        cell = ws.cell(1, col, name)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=_HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = _XLSX_COL_WIDTHS.get(name, 12)
    for r, row in enumerate(rows, 2):
        for col, name in enumerate(cols, 1):
            raw = row.get(name)
            value = "" if raw is None else str(raw)
            cell = ws.cell(r, col, value)
            if name == "trace_url" and value:
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")
            if name in ("judge_reason", "material_summary"):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            elif name in _FILL_COLUMNS:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
                cell.alignment = Alignment(horizontal="center", vertical="top")
        longest = (
            max(
                _estimate_row_height(
                    str(row.get("material_summary") or ""), _XLSX_COL_WIDTHS["material_summary"]
                ),
                _estimate_row_height(
                    str(row.get("judge_reason") or ""), _XLSX_COL_WIDTHS["judge_reason"]
                ),
            )
            if "judge_reason" in cols
            else _estimate_row_height(
                str(row.get("material_summary") or ""), _XLSX_COL_WIDTHS["material_summary"]
            )
        )
        ws.row_dimensions[r].height = longest
    ws.freeze_panes = "A2"
    wb.save(path)
