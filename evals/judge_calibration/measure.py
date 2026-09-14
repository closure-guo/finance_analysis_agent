#!/usr/bin/env python
"""judge-人工一致性校准（add-judge-human-calibration）：标注导出 + 一致性指标。

一致性指标（judge 分 vs 人工分，均为 1-5）：
- Spearman 秩相关（纯 Python 实现）
- MAE 平均绝对误差
- 方向一致率（>3 视为认可、<=3 视为不认可）
阈值配置化（env：JUDGE_MIN_SPEARMAN / JUDGE_MAX_MAE / JUDGE_MIN_DIRECTION），
低于阈值标记需校准（触发 judge prompt 修订流程，走 prompt-deploy 管线）。

标注流程：CLI（tests/scripts/judge_calibration_export.py）从 Langfuse 抽样导出
评判表（human_score 置空），人工打分后回填，再跑本模块计算一致性。

用法:
    uv run python evals/judge_calibration/measure.py --labeled path/to/labeled.jsonl
"""

from __future__ import annotations

import argparse
import json
import math
import os
import statistics
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

# 评估维度（judge 1-5 分；与 evals/run.py _JUDGE_DIMS 对齐，可配置）
DEFAULT_DIMENSIONS = (
    "report_relevance",
    "debate_quality",
    "decision_grounding",
    "consistency",
)

MIN_SPEARMAN = float(os.getenv("JUDGE_MIN_SPEARMAN", "0.5"))
MAX_MAE = float(os.getenv("JUDGE_MAX_MAE", "1.0"))
MIN_DIRECTION = float(os.getenv("JUDGE_MIN_DIRECTION", "0.7"))


@dataclass
class LabeledRow:
    """一行标注样本：judge 分（已产出）+ 人工分（待标注/已回填）。"""

    trace_id: str
    dimension: str
    judge_score: float | None = None
    human_score: float | None = None


def export_row(trace: dict[str, Any], dimension: str) -> LabeledRow:
    """从 trace 提取 judge 分：优先 trace.scores[dimension]，回退 observations 的 judge 观测。"""
    scores = trace.get("scores") or {}
    judge = None
    if isinstance(scores, dict):
        v = scores.get(dimension)
        if isinstance(v, (int, float)):
            judge = float(v)
    if judge is None:
        for o in trace.get("observations") or []:
            # traces API 的 observations 是 ID 字符串列表；内嵌对象仅
            # observations 端点返回。字符串条目跳过（调用方负责拉取 enrichment）。
            if not isinstance(o, dict):
                continue
            if str(o.get("name") or "") != "judge":
                continue
            out = o.get("output")
            if isinstance(out, dict):
                v = out.get(dimension)
                if isinstance(v, (int, float)):
                    judge = float(v)
    return LabeledRow(trace_id=str(trace.get("id") or ""), dimension=dimension, judge_score=judge)


def export_rows(
    traces: list[dict[str, Any]],
    dimensions: tuple[str, ...] = DEFAULT_DIMENSIONS,
) -> list[LabeledRow]:
    """抽样导出：每条 trace × 每个维度一行（human_score 置空待标注）。"""
    rows: list[LabeledRow] = []
    for t in traces:
        for dim in dimensions:
            rows.append(export_row(t, dim))
    return rows


def _spearman(xs: list[float], ys: list[float]) -> float | None:
    """秩相关（并列取平均秩）。"""
    n = len(xs)
    if n < 2:
        return None

    def ranks(a: list[float]) -> list[float]:
        order = sorted(range(n), key=lambda i: a[i])
        ranks_out = [0.0] * n
        i = 0
        while i < n:
            j = i
            while j + 1 < n and a[order[j + 1]] == a[order[i]]:
                j += 1
            avg_rank = (i + j) / 2 + 1
            for k in range(i, j + 1):
                ranks_out[order[k]] = avg_rank
            i = j + 1
        return ranks_out

    rx, ry = ranks(xs), ranks(ys)
    mean_x = statistics.mean(rx)
    mean_y = statistics.mean(ry)
    cov = sum((a - mean_x) * (b - mean_y) for a, b in zip(rx, ry, strict=False))
    var_x = sum((a - mean_x) ** 2 for a in rx)
    var_y = sum((b - mean_y) ** 2 for b in ry)
    if var_x == 0 or var_y == 0:
        return None
    return cov / math.sqrt(var_x * var_y)


def consistency(
    rows: list[LabeledRow],
    *,
    min_spearman: float = MIN_SPEARMAN,
    max_mae: float = MAX_MAE,
    min_direction: float = MIN_DIRECTION,
) -> dict[str, Any]:
    """一致性指标（按维度 + 整体）。低于阈值标 need_calibrate。"""
    labeled = [r for r in rows if r.judge_score is not None and r.human_score is not None]
    dims: dict[str, list[LabeledRow]] = {}
    for r in labeled:
        dims.setdefault(r.dimension, []).append(r)

    by_dim: dict[str, dict[str, Any]] = {}
    overall_judge: list[float] = []
    overall_human: list[float] = []
    for dim, group in sorted(dims.items()):
        js: list[float] = []
        hs: list[float] = []
        for r in group:
            j, h = r.judge_score, r.human_score
            if j is None or h is None:
                continue
            js.append(j)
            hs.append(h)
        overall_judge.extend(js)
        overall_human.extend(hs)
        sp = _spearman(js, hs)
        mae = round(sum(abs(j - h) for j, h in zip(js, hs, strict=False)) / len(js), 4)
        direction = round(
            sum(1 for j, h in zip(js, hs, strict=False) if (j > 3) == (h > 3)) / len(js),
            4,
        )
        by_dim[dim] = {
            "n": len(js),
            "spearman": round(sp, 4) if sp is not None else None,
            "mae": mae,
            "direction_rate": direction,
            "need_calibrate": (
                (sp is not None and sp < min_spearman) or mae > max_mae or direction < min_direction
            ),
        }

    overall: dict[str, Any] = {"n": len(labeled)}
    if labeled:
        overall["mae"] = round(
            sum(abs(j - h) for j, h in zip(overall_judge, overall_human, strict=False))
            / len(labeled),
            4,
        )
        overall["direction_rate"] = round(
            sum(1 for j, h in zip(overall_judge, overall_human, strict=False) if (j > 3) == (h > 3))
            / len(labeled),
            4,
        )
        sp = _spearman(overall_judge, overall_human)
        overall["spearman"] = round(sp, 4) if sp is not None else None
        overall["need_calibrate"] = (
            (sp is not None and sp < min_spearman)
            or overall["mae"] > max_mae
            or overall["direction_rate"] < min_direction
        )
    return {"by_dimension": by_dim, "overall": overall, "labeled_count": len(labeled)}


def render_report(result: dict[str, Any]) -> str:
    lines = [
        "# judge-人工一致性校准报告",
        "",
        f"- 生成时间: {datetime.now().isoformat(timespec='seconds')}",
        f"- 已标注样本: {result['labeled_count']}（judge 分与人工分齐备）",
        "",
        "| 维度 | n | Spearman | MAE | 方向一致率 | 需校准 |",
        "|---|---|---|---|---|---|",
    ]
    for dim, d in result["by_dimension"].items():
        lines.append(
            f"| {dim} | {d['n']} | {d['spearman'] if d['spearman'] is not None else '—'} | "
            f"{d['mae']} | {d['direction_rate']} | {'⚠️' if d['need_calibrate'] else '✓'} |"
        )
    o = result["overall"]
    if o.get("n"):
        lines.append(
            f"| **整体** | {o['n']} | {o['spearman'] if o.get('spearman') is not None else '—'} | "
            f"{o['mae']} | {o['direction_rate']} | {'⚠️' if o['need_calibrate'] else '✓'} |"
        )
    lines += [
        "",
        f"> 阈值: Spearman≥{MIN_SPEARMAN} / MAE≤{MAX_MAE} / 方向一致率≥{MIN_DIRECTION}；",
        "> 需校准 → 触发 judge prompt 修订（走 prompt-deploy 管线后重测）。",
        "",
    ]
    return "\n".join(lines)


def load_labeled_rows(path: Path) -> list[LabeledRow]:
    rows: list[LabeledRow] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        d = json.loads(line)
        rows.append(
            LabeledRow(
                trace_id=str(d.get("trace_id") or ""),
                dimension=str(d.get("dimension") or ""),
                judge_score=d.get("judge_score"),
                human_score=d.get("human_score"),
            )
        )
    return rows


def _num(v: Any) -> float | None:
    """CSV 单元格 → float；空白/非法 → None（未标注行跳过）。"""
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def load_labeled_csv(path: Path) -> list[LabeledRow]:
    """读取 CSV 标注表（export 脚本产出，含表头；human_score 空白 → None）。"""
    import csv as _csv

    rows: list[LabeledRow] = []
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        for d in _csv.DictReader(fh):
            rows.append(
                LabeledRow(
                    trace_id=str(d.get("trace_id") or ""),
                    dimension=str(d.get("dimension") or ""),
                    judge_score=_num(d.get("judge_score")),
                    human_score=_num(d.get("human_score")),
                )
            )
    return rows


def load_labeled_xlsx(path: Path) -> list[LabeledRow]:
    """读取 xlsx 标注表（export 脚本 --xlsx 产出，表头行与 CSV 同构；human_score 空 → None）。"""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    header = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {name: i for i, name in enumerate(header)}
    rows: list[LabeledRow] = []
    for vals in ws.iter_rows(min_row=2, values_only=True):
        d = {name: vals[i] if i < len(vals) else None for name, i in idx.items()}
        rows.append(
            LabeledRow(
                trace_id=str(d.get("trace_id") or ""),
                dimension=str(d.get("dimension") or ""),
                judge_score=_num(d.get("judge_score")),
                human_score=_num(d.get("human_score")),
            )
        )
    wb.close()
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description="judge-人工一致性")
    parser.add_argument(
        "--labeled", type=Path, default=None, help="标注 JSONL（每行含 human_score）"
    )
    parser.add_argument(
        "--csv", type=Path, default=None, help="标注 CSV（export 脚本产出，含表头）"
    )
    parser.add_argument(
        "--xlsx", type=Path, default=None, help="标注 xlsx（export 脚本 --xlsx 产出）"
    )
    parser.add_argument(
        "--judge-jsonl",
        type=Path,
        default=None,
        help="盲标模式下的 judge 分参考（export --blind-judge 产出）；与盲标 xlsx 合并计算一致性",
    )
    parser.add_argument("--out", type=Path, default=Path("reports/judge-calibration-report.md"))
    args = parser.parse_args()
    given = [flag for flag in ("--labeled", "--csv", "--xlsx") if getattr(args, flag.lstrip("-"))]
    if len(given) != 1:
        parser.error(f"必须且只能指定 {'/'.join(given)} 之一：--labeled、--csv、--xlsx")

    if args.xlsx:
        rows = load_labeled_xlsx(args.xlsx)
        if args.judge_jsonl:
            _judge_map: dict[tuple[str, str], float] = {}
            for line in args.judge_jsonl.read_text(encoding="utf-8").splitlines():
                if not line.strip():
                    continue
                d = json.loads(line)
                _judge_map[(str(d["trace_id"]), str(d["dimension"]))] = float(d["judge_score"])
            for r in rows:
                j = _judge_map.get((r.trace_id, r.dimension))
                if j is not None:
                    r.judge_score = j
    elif args.csv:
        rows = load_labeled_csv(args.csv)
    else:
        rows = load_labeled_rows(args.labeled)

    result = consistency(rows)
    text = render_report(result)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(text, encoding="utf-8")
    print(text)


if __name__ == "__main__":
    main()
